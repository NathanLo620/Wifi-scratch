#!/usr/bin/env python3
"""
Rebuild combo_percentile_summary from raw histogram CSVs and export to Excel.

Outputs:
  combo_percentile_summary_1Mbps.csv   (raw data)
  combo_percentile_analysis.xlsx       (multi-sheet Excel with rankings)
"""

import csv
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule

# ── Paths ──────────────────────────────────────────────────────────────
BASE      = Path(__file__).resolve().parent
N_STA     = 30
DATA_RATE = "1Mbps"
N_PEDCA_LIST = [5, 15, 30]
CWDS_VALUES  = [0, 1]
QSRC_VALUES  = list(range(0, 6))   # 0..5
PSRC_VALUES  = [1, 2, 3]

# ── Histogram loader ───────────────────────────────────────────────────

def load_histogram(csv_path: Path):
    rows = []
    with csv_path.open("r", newline="") as f:
        for row in csv.DictReader(f):
            rows.append({
                "start": float(row["bin_start_us"]),
                "end":   float(row["bin_end_us"]),
                "prob":  float(row["probability"]),
            })
    if not rows:
        return None
    widths = sorted(r["end"] - r["start"] for r in rows if r["end"] > r["start"])
    bw = widths[len(widths) // 2]
    min_s = min(r["start"] for r in rows)
    max_e = max(r["end"]   for r in rows)
    lu = {round(r["start"] / bw) * bw: r["prob"] for r in rows}
    mids, probs = [], []
    cur, eps = min_s, bw * 1e-6
    while cur < max_e - eps:
        mids.append(cur + 0.5 * bw)
        probs.append(lu.get(round(cur / bw) * bw, 0.0))
        cur += bw
    return mids, probs


def compute_percentiles(mids, probs, pcts=(0.5, 0.95, 0.99)):
    total = sum(probs)
    if total <= 0:
        return {p: None for p in pcts}
    result, running = {}, 0.0
    targets = sorted(pcts)
    ti = 0
    for m, p in zip(mids, probs):
        running += p
        while ti < len(targets) and running / total >= targets[ti]:
            result[targets[ti]] = m
            ti += 1
        if ti >= len(targets):
            break
    for pct in pcts:
        if pct not in result:
            result[pct] = mids[-1] if mids else None
    return result


# ── Collect all records ────────────────────────────────────────────────

records = []

# EDCA-only baseline
edca_csv = BASE / "edca_only" / f"edca_only_p00_vo_delay_pdf_nSta{N_STA}_{DATA_RATE}.csv"
edca_pcts = {}
if edca_csv.exists():
    r = load_histogram(edca_csv)
    if r:
        edca_pcts = compute_percentiles(r[0], r[1])

# All combos
for cwds in CWDS_VALUES:
    for qsrc in QSRC_VALUES:
        for psrc in PSRC_VALUES:
            tag = f"c{cwds}_q{qsrc:02d}_s{psrc:02d}"
            d   = BASE / tag
            if not d.exists():
                continue
            for n_pedca in N_PEDCA_LIST:
                paths = {
                    "all":    d / f"{tag}_p{n_pedca:02d}_vo_delay_pdf_nSta{N_STA}_{DATA_RATE}.csv",
                    "pedca":  d / f"{tag}_p{n_pedca:02d}_pedca_sta_delay_pdf_nSta{N_STA}_{DATA_RATE}.csv",
                    "legacy": d / f"{tag}_p{n_pedca:02d}_legacy_sta_delay_pdf_nSta{N_STA}_{DATA_RATE}.csv",
                }
                for sta_type, path in paths.items():
                    if not path.exists():
                        continue
                    h = load_histogram(path)
                    if h is None:
                        continue
                    pcts = compute_percentiles(h[0], h[1])
                    if pcts[0.5] is None:
                        continue
                    records.append({
                        "CWds": cwds, "QSRC": qsrc, "PSRC": psrc,
                        "nPedca": n_pedca, "sta_type": sta_type,
                        "P50_us": pcts[0.5],
                        "P95_us": pcts[0.95],
                        "P99_us": pcts[0.99],
                    })

# ── Write raw CSV ──────────────────────────────────────────────────────

csv_out = BASE / f"combo_percentile_summary_{DATA_RATE}.csv"
with open(csv_out, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["CWds", "QSRC", "PSRC", "nPedca", "delay_type",
                "P50_us", "P95_us", "P99_us"])
    for r in sorted(records, key=lambda x: (x["sta_type"], x["nPedca"],
                                              x["CWds"], x["QSRC"], x["PSRC"])):
        w.writerow([r["CWds"], r["QSRC"], r["PSRC"], r["nPedca"], r["sta_type"],
                    f"{r['P50_us']:.2f}", f"{r['P95_us']:.2f}", f"{r['P99_us']:.2f}"])
print(f"  ✔ {csv_out.name}  ({len(records)} rows)")


# ── Excel helpers ──────────────────────────────────────────────────────

def hex_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

HEADER_FILL   = hex_fill("1C2128")
TITLE_FILL    = hex_fill("0D1117")
GOLD_FILL     = hex_fill("3D2B00")
BLUE_FILL     = hex_fill("0D1F3C")
GREEN_FILL    = hex_fill("0D2B1A")
ORANGE_FILL   = hex_fill("2E1A00")
PURPLE_FILL   = hex_fill("1F1240")
BEST_GOLD     = hex_fill("4D3500")
BEST_BLUE     = hex_fill("0A2A50")
BEST_GREEN    = hex_fill("0A2E1A")
ZEBRA_DARK    = hex_fill("0F1117")
ZEBRA_LIGHT   = hex_fill("161B22")

WHITE    = Font(color="F0F6FC", bold=False)
BOLD_W   = Font(color="F0F6FC", bold=True)
GREY_F   = Font(color="8B949E")
GOLD_F   = Font(color="E3B341", bold=True)
BLUE_F   = Font(color="79C0FF", bold=True)
GREEN_F  = Font(color="56D364", bold=True)
ORANGE_F = Font(color="FFA657", bold=True)

CENTER = Alignment(horizontal="center", vertical="center")
RIGHT  = Alignment(horizontal="right",  vertical="center")
LEFT   = Alignment(horizontal="left",   vertical="center")


def thin_border():
    s = Side(style="thin", color="30363D")
    return Border(left=s, right=s, top=s, bottom=s)


def set_header_row(ws, row_idx, cols, fills=None, fonts=None):
    for ci, val in enumerate(cols, 1):
        c = ws.cell(row=row_idx, column=ci, value=val)
        c.fill  = fills[ci-1] if fills else HEADER_FILL
        c.font  = fonts[ci-1] if fonts else BOLD_W
        c.alignment = CENTER
        c.border = thin_border()


def write_data_row(ws, row_idx, values, is_best_p50=False, is_best_p95=False,
                   is_best_p99=False, zebra=False):
    fill = ZEBRA_LIGHT if zebra else ZEBRA_DARK
    if is_best_p50 and is_best_p95 and is_best_p99:
        fill = PURPLE_FILL
    elif is_best_p50:
        fill = BLUE_FILL
    elif is_best_p95:
        fill = ORANGE_FILL
    elif is_best_p99:
        fill = GREEN_FILL

    for ci, val in enumerate(values, 1):
        c = ws.cell(row=row_idx, column=ci, value=val)
        c.fill   = fill
        c.font   = WHITE
        c.border = thin_border()
        if isinstance(val, (int, float)):
            c.alignment = RIGHT
            c.number_format = "#,##0.0"
        else:
            c.alignment = LEFT


def freeze_and_width(ws, freeze_cell, col_widths):
    ws.freeze_panes = freeze_cell
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── Build Excel ────────────────────────────────────────────────────────

wb = openpyxl.Workbook()
wb.remove(wb.active)   # remove default sheet

STA_TYPES = [
    ("all",    "All STAs",      GOLD_FILL,   GOLD_F),
    ("pedca",  "P-EDCA STAs",   BLUE_FILL,   BLUE_F),
    ("legacy", "Legacy STAs",   GREEN_FILL,  GREEN_F),
]

METRIC_KEYS = ["P50_us", "P95_us", "P99_us"]
METRIC_LABELS = {"P50_us": "Median (P50) µs",
                 "P95_us": "P95 µs",
                 "P99_us": "P99 µs"}

baseline_vals = {
    "P50_us": edca_pcts.get(0.5),
    "P95_us": edca_pcts.get(0.95),
    "P99_us": edca_pcts.get(0.99),
}

# ──────────────────────────────────────────────────────────────────────
# Sheet 1-3: per-STA-type full ranked tables (one sheet per type)
# ──────────────────────────────────────────────────────────────────────

for sta_type, type_label, type_fill, type_font in STA_TYPES:
    ws = wb.create_sheet(title=type_label)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1F6FEB" if sta_type == "all" else (
        "3B82F6" if sta_type == "pedca" else "22C55E")

    # Title
    ws.merge_cells("A1:J1")
    tc = ws["A1"]
    tc.value = f"{type_label} — VO Delay Percentile Sweep  |  nSta={N_STA}  dataRate={DATA_RATE}  CWds×QSRC×PSRC×nPedca"
    tc.fill, tc.font = TITLE_FILL, Font(color="F0F6FC", bold=True, size=12)
    tc.alignment = LEFT
    ws.row_dimensions[1].height = 22

    # Baseline row
    ws.merge_cells("A2:D2")
    ws["A2"].value = f"EDCA-only baseline (nPedca=0)"
    ws["A2"].fill  = hex_fill("161B22")
    ws["A2"].font  = GREY_F
    ws["A2"].alignment = LEFT
    for ci, mk in enumerate(METRIC_KEYS, 5):
        c = ws.cell(row=2, column=ci)
        c.value = baseline_vals.get(mk)
        c.fill  = hex_fill("161B22")
        c.font  = Font(color="E3B341", bold=True)
        c.alignment = RIGHT
        c.number_format = "#,##0.0"
    ws.row_dimensions[2].height = 16

    # Column headers
    headers = ["CWds", "QSRC", "PSRC", "nPedca",
               "Median (P50) µs", "P95 µs", "P99 µs",
               "ΔP50 vs EDCA%", "ΔP95 vs EDCA%", "ΔP99 vs EDCA%"]
    set_header_row(ws, 3, headers)
    ws.row_dimensions[3].height = 18

    # Data: sort by (nPedca, P50)
    subset = [r for r in records if r["sta_type"] == sta_type]
    subset_sorted = sorted(subset, key=lambda x: (x["nPedca"], x["P50_us"]))

    # Per-nPedca: find best values for highlighting
    best_map = {}
    for n in N_PEDCA_LIST:
        s = [r for r in subset if r["nPedca"] == n]
        if not s:
            continue
        best_map[n] = {
            "P50": min(s, key=lambda x: x["P50_us"])["P50_us"],
            "P95": min(s, key=lambda x: x["P95_us"])["P95_us"],
            "P99": min(s, key=lambda x: x["P99_us"])["P99_us"],
        }

    row = 4
    prev_npedca = None
    for i, r in enumerate(subset_sorted):
        n = r["nPedca"]
        if n != prev_npedca:
            # nPedca separator
            ws.merge_cells(f"A{row}:J{row}")
            sc = ws[f"A{row}"]
            sc.value = f"── nPedca = {n}/{N_STA} ──"
            sc.fill  = hex_fill("1C2128")
            sc.font  = Font(color="79C0FF", bold=True, size=10)
            sc.alignment = LEFT
            ws.row_dimensions[row].height = 15
            row += 1
            prev_npedca = n

        bm = best_map.get(n, {})
        is_best_p50 = (r["P50_us"] == bm.get("P50"))
        is_best_p95 = (r["P95_us"] == bm.get("P95"))
        is_best_p99 = (r["P99_us"] == bm.get("P99"))

        def delta_pct(val, base):
            if base and base > 0:
                return round((val - base) / base * 100, 2)
            return None

        d50 = delta_pct(r["P50_us"], baseline_vals["P50_us"])
        d95 = delta_pct(r["P95_us"], baseline_vals["P95_us"])
        d99 = delta_pct(r["P99_us"], baseline_vals["P99_us"])

        values = [r["CWds"], r["QSRC"], r["PSRC"], r["nPedca"],
                  r["P50_us"], r["P95_us"], r["P99_us"], d50, d95, d99]
        write_data_row(ws, row, values,
                       is_best_p50=is_best_p50,
                       is_best_p95=is_best_p95,
                       is_best_p99=is_best_p99,
                       zebra=(i % 2 == 0))

        # Badge text for best cells
        if is_best_p50:
            ws.cell(row=row, column=5).font = Font(color="79C0FF", bold=True)
        if is_best_p95:
            ws.cell(row=row, column=6).font = Font(color="FFA657", bold=True)
        if is_best_p99:
            ws.cell(row=row, column=7).font = Font(color="56D364", bold=True)

        ws.row_dimensions[row].height = 15
        row += 1

    freeze_and_width(ws, "A4",
                     [7, 7, 7, 8, 16, 13, 13, 17, 17, 17])


# ──────────────────────────────────────────────────────────────────────
# Sheet 4: Best Parameters Summary (one row per sta_type × metric × nPedca)
# ──────────────────────────────────────────────────────────────────────

ws_best = wb.create_sheet(title="Best Params Summary")
ws_best.sheet_view.showGridLines = False

ws_best.merge_cells("A1:K1")
c = ws_best["A1"]
c.value = "Best Parameter Combinations — Minimum Delay per STA Type × Metric × nPedca"
c.fill, c.font = TITLE_FILL, Font(color="F0F6FC", bold=True, size=12)
c.alignment = LEFT
ws_best.row_dimensions[1].height = 22

headers_best = ["STA Type", "nPedca", "Metric",
                "Best CWds", "Best QSRC", "Best PSRC",
                "Best Value (µs)", "EDCA Baseline (µs)", "Δ vs EDCA (%)",
                "Rank 2 Value (µs)", "Rank 2 Params"]
set_header_row(ws_best, 2, headers_best)
ws_best.row_dimensions[2].height = 18

row = 3
for sta_type, type_label, type_fill, type_font in STA_TYPES:
    for n in N_PEDCA_LIST:
        subset = [r for r in records
                  if r["sta_type"] == sta_type and r["nPedca"] == n]
        if not subset:
            continue

        for metric_key, metric_label in [
            ("P50_us", "Median (P50)"),
            ("P95_us", "P95"),
            ("P99_us", "P99"),
        ]:
            s = sorted(subset, key=lambda x: x[metric_key])
            best = s[0]
            rank2 = s[1] if len(s) > 1 else None
            base_val = baseline_vals.get(metric_key)
            delta = round((best[metric_key] - base_val) / base_val * 100, 2) if base_val else None

            values = [
                type_label, n, metric_label,
                best["CWds"], best["QSRC"], best["PSRC"],
                round(best[metric_key], 1),
                round(base_val, 1) if base_val else None,
                delta,
                round(rank2[metric_key], 1) if rank2 else None,
                f"c{rank2['CWds']}·q{rank2['QSRC']}·s{rank2['PSRC']}" if rank2 else "",
            ]

            for ci, val in enumerate(values, 1):
                c = ws_best.cell(row=row, column=ci, value=val)
                c.fill   = type_fill
                c.font   = WHITE
                c.border = thin_border()
                if isinstance(val, (int, float)):
                    c.alignment = RIGHT
                    if ci == 9:   # delta %
                        c.number_format = '+0.00%;-0.00%' if val else "#,##0.0"
                        c.value = (delta / 100.0) if delta is not None else None
                        c.number_format = '+0.0%;-0.0%'
                        c.font = Font(
                            color="56D364" if (delta or 0) < 0 else "FFA657",
                            bold=True)
                    else:
                        c.number_format = "#,##0.0"
                else:
                    c.alignment = LEFT

            # Colour-code the "Best Value" cell
            bv_cell = ws_best.cell(row=row, column=7)
            if metric_key == "P50_us":
                bv_cell.font = Font(color="79C0FF", bold=True)
            elif metric_key == "P95_us":
                bv_cell.font = Font(color="FFA657", bold=True)
            else:
                bv_cell.font = Font(color="56D364", bold=True)

            ws_best.row_dimensions[row].height = 15
            row += 1

        # blank separator between nPedca groups within same sta_type
        for ci in range(1, 12):
            c = ws_best.cell(row=row, column=ci)
            c.fill   = hex_fill("0D1117")
            c.border = thin_border()
        ws_best.row_dimensions[row].height = 5
        row += 1

freeze_and_width(ws_best, "A3",
                 [14, 8, 13, 10, 10, 10, 17, 18, 14, 18, 16])


# ──────────────────────────────────────────────────────────────────────
# Sheet 5: Heatmap — P95 vs (QSRC, PSRC) for each (STA type, nPedca, CWds)
# ──────────────────────────────────────────────────────────────────────

ws_hm = wb.create_sheet(title="P95 Heatmap")
ws_hm.sheet_view.showGridLines = False

ws_hm.merge_cells("A1:Z1")
c = ws_hm["A1"]
c.value = "P95 Delay Heatmap — lower is better  (CWds=0 left, CWds=1 right)"
c.fill, c.font = TITLE_FILL, Font(color="F0F6FC", bold=True, size=12)
c.alignment = LEFT
ws_hm.row_dimensions[1].height = 22

current_row = 2
for sta_type, type_label, type_fill, _ in STA_TYPES:
    for n in N_PEDCA_LIST:
        subset = [r for r in records
                  if r["sta_type"] == sta_type and r["nPedca"] == n]
        if not subset:
            continue

        # Section header
        ws_hm.merge_cells(f"A{current_row}:Z{current_row}")
        sc = ws_hm[f"A{current_row}"]
        sc.value = f"{type_label}  |  nPedca={n}/{N_STA}"
        sc.fill, sc.font = type_fill, Font(color="F0F6FC", bold=True)
        sc.alignment = LEFT
        ws_hm.row_dimensions[current_row].height = 16
        current_row += 1

        # Column layout: [label | CWds=0 PSRC1..3 | CWds=1 PSRC1..3]  × QSRC rows
        col_start = 1
        # Header row: PSRC across, CWds groups
        hdr_vals = ["QSRC \\ PSRC"] + \
                   [f"CWds=0 P{p}" for p in PSRC_VALUES] + \
                   [f"CWds=1 P{p}" for p in PSRC_VALUES]
        for ci, v in enumerate(hdr_vals, col_start):
            c = ws_hm.cell(row=current_row, column=ci, value=v)
            c.fill, c.font = HEADER_FILL, BOLD_W
            c.alignment = CENTER
            c.border = thin_border()
        ws_hm.row_dimensions[current_row].height = 16
        current_row += 1

        # Build lookup {(cwds, qsrc, psrc): p95}
        lu = {(r["CWds"], r["QSRC"], r["PSRC"]): r["P95_us"] for r in subset}
        all_p95 = [v for v in lu.values() if v is not None]
        min_p95 = min(all_p95) if all_p95 else 0
        max_p95 = max(all_p95) if all_p95 else 1

        for qsrc in QSRC_VALUES:
            row_vals = [f"QSRC={qsrc}"]
            for cwds in CWDS_VALUES:
                for psrc in PSRC_VALUES:
                    row_vals.append(lu.get((cwds, qsrc, psrc)))

            for ci, val in enumerate(row_vals, col_start):
                c = ws_hm.cell(row=current_row, column=ci, value=val)
                c.border = thin_border()
                if isinstance(val, (int, float)) and val is not None:
                    c.alignment = RIGHT
                    c.number_format = "#,##0.0"
                    # Colour: green (best) → red (worst)
                    ratio = (val - min_p95) / (max_p95 - min_p95 + 1)
                    r_c = int(30 + ratio * 180)
                    g_c = int(150 - ratio * 120)
                    b_c = 30
                    hex_col = f"{r_c:02X}{g_c:02X}{b_c:02X}"
                    c.fill = hex_fill(hex_col)
                    c.font = Font(color="F0F6FC" if ratio > 0.5 else "0D1117",
                                  bold=(val == min_p95))
                    if val == min_p95:
                        c.font = Font(color="F0F6FC", bold=True)
                else:
                    c.fill = hex_fill("161B22")
                    c.font = BOLD_W
                    c.alignment = LEFT
            ws_hm.row_dimensions[current_row].height = 14
            current_row += 1

        current_row += 1   # blank between groups

ws_hm.column_dimensions["A"].width = 12
for ci in range(2, 2 + len(CWDS_VALUES) * len(PSRC_VALUES)):
    ws_hm.column_dimensions[get_column_letter(ci)].width = 13
ws_hm.freeze_panes = "B3"


# ── Save ───────────────────────────────────────────────────────────────

xl_out = BASE / "combo_percentile_analysis.xlsx"
wb.save(xl_out)
print(f"  ✔ {xl_out.name}  ({xl_out.stat().st_size:,} bytes)")
print(f"\n  Summary: {len(records)} combo×nPedca×sta_type records")
print(f"  EDCA baseline: P50={edca_pcts.get(0.5):.0f}µs  P95={edca_pcts.get(0.95):.0f}µs  P99={edca_pcts.get(0.99):.0f}µs")

# Print quick best-params table to stdout
print(f"\n{'='*85}")
print(f"  Best parameters per (STA type × metric × nPedca)")
print(f"{'='*85}")
print(f"  {'STA':8} {'nPedca':7} {'Metric':10} {'CWds':5} {'QSRC':5} {'PSRC':5}  {'Best (µs)':>10}  {'EDCA (µs)':>10}  {'Δ%':>8}")
print(f"  {'-'*80}")
for sta_type, type_label, _, _ in STA_TYPES:
    for n in N_PEDCA_LIST:
        subset = [r for r in records if r["sta_type"] == sta_type and r["nPedca"] == n]
        if not subset:
            continue
        for mk, ml in [("P50_us","P50"),("P95_us","P95"),("P99_us","P99")]:
            best = min(subset, key=lambda x: x[mk])
            base = baseline_vals.get(mk, 0) or 1
            delta = (best[mk] - base) / base * 100
            sign = "+" if delta >= 0 else ""
            print(f"  {type_label:8} {n:7} {ml:10} "
                  f"{best['CWds']:5} {best['QSRC']:5} {best['PSRC']:5}  "
                  f"{best[mk]:>10.0f}  {base:>10.0f}  {sign}{delta:>7.1f}%")
        print(f"  {'·'*80}")
