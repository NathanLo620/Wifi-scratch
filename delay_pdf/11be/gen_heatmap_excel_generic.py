#!/usr/bin/env python3
"""
P50/P95/P99 delay HEATMAP workbook, colour-scaled (green=low/best -> red=high/worst).

Generic version of gen_heatmap_excel.py: takes the combo_percentile_summary
CSV and output xlsx path as CLI args, so it can be pointed at either the
mono-DS data (top-level combo_percentile_summary_<rate>.csv) or the dual-DS
data (dual-DS/combo_percentile_summary_<rate>.csv) for any traffic model.

Usage: gen_heatmap_excel_generic.py <summary.csv> <out.xlsx> [--rate 1Mbps]
"""

import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

RATE = "1Mbps"
if "--rate" in sys.argv:
    RATE = sys.argv[sys.argv.index("--rate") + 1]

SUMMARY = Path(sys.argv[1])
OUT = Path(sys.argv[2])

QSRC = [0, 1, 2, 3, 4, 5]
PSRC = [1, 2, 3]
CWDS = [0, 1]
STA_TYPES = [("vo", "All STAs"), ("pedca", "P-EDCA STAs"), ("legacy", "Legacy STAs")]
N_FOR = {"vo": [5, 15, 30], "pedca": [5, 15, 30], "legacy": [5, 15]}

P50, P95, P99 = {}, {}, {}
for r in csv.DictReader(open(SUMMARY)):
    key = (r["delay_type"], int(r["CWds"]), int(r["QSRC"]),
           int(r["PSRC"]), int(r["nPedca"]))
    P50[key] = float(r["P50_us"])
    P95[key] = float(r["P95_us"])
    P99[key] = float(r["P99_us"])

def fill(h):
    return PatternFill("solid", fgColor=h)

TITLE_FILL = fill("0D1117")
SEC_FILL   = fill("1C2128")
HDR_FILL   = fill("21262D")
WHITE  = Font(color="F0F6FC")
BOLD_W = Font(color="F0F6FC", bold=True)
BLUE   = Font(color="79C0FF", bold=True)
GOLD   = Font(color="E3B341", bold=True)
DARK   = Font(color="0D1117", bold=True)
CENTER = Alignment("center", "center")
LEFT   = Alignment("left", "center")

def bd():
    s = Side(style="thin", color="30363D")
    return Border(s, s, s, s)

def color_scale():
    return ColorScaleRule(
        start_type="min",  start_color="2DA44E",
        mid_type="percentile", mid_value=50, mid_color="E3B341",
        end_type="max",    end_color="F85149",
    )

wb = openpyxl.Workbook()
wb.remove(wb.active)

def put(ws, r, c, v, f=WHITE, bg=None, al=CENTER, num=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = f
    if bg: x.fill = bg
    x.alignment = al
    x.border = bd()
    if num: x.number_format = num
    return x

METRICS = [("P50", P50), ("P95", P95), ("P99", P99)]

for sta_type, label in STA_TYPES:
    ws = wb.create_sheet(title=label)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = {"vo": "E3B341", "pedca": "1F6FEB",
                                    "legacy": "2DA44E"}[sta_type]

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = (f"{label} — delay heatmap (µs)   |   "
               f"rows = QSRC, cols = PSRC   |   green = lower (better), red = higher")
    t.fill, t.font, t.alignment = TITLE_FILL, Font(color="F0F6FC", bold=True, size=12), LEFT
    ws.row_dimensions[1].height = 24

    row = 3
    for n in N_FOR[sta_type]:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        s = ws.cell(row=row, column=1,
                    value=f"▼ nPedca = {n} / 30        (left block CWds=0   ·   right block CWds=1)")
        s.fill, s.font, s.alignment = SEC_FILL, GOLD, LEFT
        row += 1

        for metric_label, metric_dict in METRICS:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
            m = ws.cell(row=row, column=1, value=f"   {metric_label}")
            m.fill, m.font, m.alignment = HDR_FILL, BLUE, LEFT
            row += 1

            block_cols = {0: 1, 1: 6}
            for cwds in CWDS:
                c0 = block_cols[cwds]
                put(ws, row, c0, f"CWds={cwds}", GOLD, HDR_FILL)
                for j, ps in enumerate(PSRC, 1):
                    put(ws, row, c0 + j, f"PSRC={ps}", BOLD_W, HDR_FILL)
                data_top = row + 1
                for i, q in enumerate(QSRC):
                    rr = data_top + i
                    put(ws, rr, c0, f"QSRC={q}", BOLD_W, HDR_FILL, LEFT)
                    for j, ps in enumerate(PSRC, 1):
                        v = metric_dict.get((sta_type, cwds, q, ps, n))
                        put(ws, rr, c0 + j, round(v, 0) if v is not None else None,
                            DARK, None, CENTER, "#,##0")
                rng = (f"{get_column_letter(c0+1)}{data_top}:"
                       f"{get_column_letter(c0+3)}{data_top+len(QSRC)-1}")
                ws.conditional_formatting.add(rng, color_scale())
            row = data_top + len(QSRC) + 1
        row += 1

    for col in [1, 6]:
        ws.column_dimensions[get_column_letter(col)].width = 11
    for col in [2, 3, 4, 7, 8, 9]:
        ws.column_dimensions[get_column_letter(col)].width = 10
    ws.column_dimensions["E"].width = 3
    ws.freeze_panes = "A3"

wb.save(OUT)
print(f"wrote {OUT}")
