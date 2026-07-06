#!/usr/bin/env python3
"""
P99 VO-delay HEATMAP workbook.

For each viewpoint (All / P-EDCA / Legacy STAs) and each nPedca regime,
draw a QSRC(rows) x PSRC(cols) grid for CWds=0 and CWds=1 side-by-side,
coloured by a green(low)->red(high) colour scale so the best region is
visually obvious.  Built from combo_percentile_summary_1Mbps.csv.

Output: param_p99_heatmap_1Mbps.xlsx
"""

import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

BASE = Path("/home/wmnlab/Desktop/ns-3.45/scratch/delay_pdf/fix_nsta30_CwdsxQSRCxPSRC_sweep_MMPP")
RATE = "1Mbps"
SUMMARY = BASE / f"combo_percentile_summary_{RATE}.csv"

QSRC = [0, 1, 2, 3, 4, 5]
PSRC = [1, 2, 3]
CWDS = [0, 1]
STA_TYPES = [("all", "All STAs"), ("pedca", "P-EDCA STAs"), ("legacy", "Legacy STAs")]

# nPedca regimes present per sta_type
N_FOR = {"all": [5, 15, 30], "pedca": [5, 15, 30], "legacy": [5, 15]}

# index: (sta_type, cwds, qsrc, psrc, nPedca) -> P99
P99 = {}
for r in csv.DictReader(open(SUMMARY)):
    key = (r["delay_type"], int(r["CWds"]), int(r["QSRC"]),
           int(r["PSRC"]), int(r["nPedca"]))
    P99[key] = float(r["P99_us"])

# ── styling ──────────────────────────────────────────────────────────
def fill(h):
    return PatternFill("solid", fgColor=h)

TITLE_FILL = fill("0D1117")
SEC_FILL   = fill("1C2128")
HDR_FILL   = fill("21262D")
WHITE  = Font(color="F0F6FC")
BOLD_W = Font(color="F0F6FC", bold=True)
GREY   = Font(color="8B949E")
BLUE   = Font(color="79C0FF", bold=True)
GOLD   = Font(color="E3B341", bold=True)
DARK   = Font(color="0D1117", bold=True)     # readable on bright heat cells
CENTER = Alignment("center", "center")
LEFT   = Alignment("left", "center")

def bd():
    s = Side(style="thin", color="30363D")
    return Border(s, s, s, s)

# Green(low/best) -> Yellow -> Red(high/worst)
def color_scale():
    return ColorScaleRule(
        start_type="min",  start_color="2DA44E",
        mid_type="percentile", mid_value=50, mid_color="E3B341",
        end_type="max",    end_color="F85149",
    )

wb = openpyxl.Workbook()
wb.remove(wb.active)

def put(ws, r, c, v, f=WHITE, bg=None, al=CENTER, num=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = f
    if bg: x.fill = bg
    x.alignment = al
    x.border = bd()
    if num: x.number_format = num
    return x

for sta_type, label in STA_TYPES:
    ws = wb.create_sheet(title=label)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = {"all": "E3B341", "pedca": "1F6FEB",
                                    "legacy": "2DA44E"}[sta_type]

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = (f"{label} — P99 VO delay heatmap (µs)   |   "
               f"rows = QSRC, cols = PSRC   |   green = lower (better), red = higher")
    t.fill, t.font, t.alignment = TITLE_FILL, Font(color="F0F6FC", bold=True, size=12), LEFT
    ws.row_dimensions[1].height = 24

    row = 3
    for n in N_FOR[sta_type]:
        # section header spanning both CWds blocks
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        s = ws.cell(row=row, column=1,
                    value=f"▼ nPedca = {n} / 30        (left block CWds=0   ·   right block CWds=1)")
        s.fill, s.font, s.alignment = SEC_FILL, BLUE, LEFT
        row += 1

        # two blocks: CWds=0 at cols 1-4, CWds=1 at cols 6-9
        block_cols = {0: 1, 1: 6}   # starting column of the QSRC-label col
        for cwds in CWDS:
            c0 = block_cols[cwds]
            # block label + PSRC header
            put(ws, row, c0, f"CWds={cwds}", GOLD, HDR_FILL)
            for j, ps in enumerate(PSRC, 1):
                put(ws, row, c0 + j, f"PSRC={ps}", BOLD_W, HDR_FILL)
            data_top = row + 1
            for i, q in enumerate(QSRC):
                rr = data_top + i
                put(ws, rr, c0, f"QSRC={q}", BOLD_W, HDR_FILL, LEFT)
                for j, ps in enumerate(PSRC, 1):
                    v = P99.get((sta_type, cwds, q, ps, n))
                    put(ws, rr, c0 + j, round(v, 0) if v is not None else None,
                        DARK, None, CENTER, "#,##0")
            # color scale over this block's data cells
            rng = (f"{get_column_letter(c0+1)}{data_top}:"
                   f"{get_column_letter(c0+3)}{data_top+len(QSRC)-1}")
            ws.conditional_formatting.add(rng, color_scale())
        row = data_top + len(QSRC) + 1   # gap before next regime

    # widths
    for col in [1, 6]:
        ws.column_dimensions[get_column_letter(col)].width = 11
    for col in [2, 3, 4, 7, 8, 9]:
        ws.column_dimensions[get_column_letter(col)].width = 10
    ws.column_dimensions["E"].width = 3
    ws.freeze_panes = "A3"

out = BASE / f"param_p99_heatmap_{RATE}.xlsx"
wb.save(out)
print(f"✔ wrote {out.name}")
