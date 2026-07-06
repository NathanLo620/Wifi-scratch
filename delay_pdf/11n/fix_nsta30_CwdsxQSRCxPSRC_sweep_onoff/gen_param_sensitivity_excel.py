#!/usr/bin/env python3
"""
Parameter-sensitivity analysis for P99 VO delay.

Question answered:
  For each parameter (CWds, QSRC, PSRC) and each viewpoint
  (P-EDCA STAs / Legacy STAs / All STAs), which direction of the
  parameter minimises the P99 delay?  e.g. "smaller QSRC -> smaller delay".

Method:
  Read combo_percentile_summary_1Mbps.csv (one row per
  CWds x QSRC x PSRC x nPedca x sta_type).  For each (sta_type, parameter)
  compute the MARGINAL mean P99 at every level of that parameter
  (averaging over the other two parameters and over nPedca), the Pearson
  correlation between level and P99, the best (min-mean) level, and a
  plain-language trend direction.

Output:
  param_sensitivity_p99_1Mbps.xlsx
"""

import csv
from pathlib import Path
from statistics import mean

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path("/home/wmnlab/Desktop/ns-3.45/scratch/delay_pdf/fix_nsta30_CwdsxQSRCxPSRC_sweep_onoff")
DATA_RATE = "1Mbps"
SUMMARY = BASE / f"combo_percentile_summary_{DATA_RATE}.csv"

PARAMS = ["CWds", "QSRC", "PSRC"]
PARAM_LEVELS = {"CWds": [0, 1], "QSRC": [0, 1, 2, 3, 4, 5], "PSRC": [1, 2, 3]}
N_PEDCA_LIST = [5, 15, 30]
METRIC = "P99_us"

STA_TYPES = [
    ("all",    "All STAs"),
    ("pedca",  "P-EDCA STAs"),
    ("legacy", "Legacy STAs"),
]

# ── Load ────────────────────────────────────────────────────────────────
records = []
for r in csv.DictReader(open(SUMMARY)):
    records.append({
        "CWds": int(r["CWds"]), "QSRC": int(r["QSRC"]), "PSRC": int(r["PSRC"]),
        "nPedca": int(r["nPedca"]), "sta_type": r["delay_type"],
        "P50_us": float(r["P50_us"]), "P95_us": float(r["P95_us"]),
        "P99_us": float(r["P99_us"]),
    })


def pearson(xs, ys):
    n = len(xs)
    if n < 2:
        return None
    mx, my = mean(xs), mean(ys)
    sxy = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    sxx = sum((x - mx) ** 2 for x in xs)
    syy = sum((y - my) ** 2 for y in ys)
    if sxx == 0 or syy == 0:
        return None
    return sxy / (sxx * syy) ** 0.5


def trend_text(corr):
    if corr is None:
        return "n/a"
    if abs(corr) < 0.15:
        return "no clear trend (≈flat)"
    direction = "larger" if corr > 0 else "smaller"
    strength = "strongly" if abs(corr) >= 0.5 else "mildly"
    return f"{strength}: {direction} param → larger delay"


# ── Styling ───────────────────────────────────────────────────────────
def fill(h):
    return PatternFill("solid", fgColor=h)

TITLE_FILL = fill("0D1117")
HEAD_FILL  = fill("1C2128")
ZEBRA_A    = fill("0F1117")
ZEBRA_B    = fill("161B22")
BEST_FILL  = fill("0D2B1A")
TYPE_FILL  = {"all": fill("3D2B00"), "pedca": fill("0D1F3C"), "legacy": fill("0D2B1A")}

WHITE  = Font(color="F0F6FC")
BOLD_W = Font(color="F0F6FC", bold=True)
GREY   = Font(color="8B949E")
GREEN  = Font(color="56D364", bold=True)
BLUE   = Font(color="79C0FF", bold=True)
GOLD   = Font(color="E3B341", bold=True)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT   = Alignment(horizontal="left", vertical="center")
RIGHT  = Alignment(horizontal="right", vertical="center")


def border():
    s = Side(style="thin", color="30363D")
    return Border(left=s, right=s, top=s, bottom=s)


def cell(ws, r, c, v, f=WHITE, bg=None, al=LEFT, num=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font = f
    if bg:
        x.fill = bg
    x.alignment = al
    x.border = border()
    if num:
        x.number_format = num
    return x


wb = openpyxl.Workbook()
wb.remove(wb.active)

# ── Conclusion sheet built up while we go ───────────────────────────────
conclusion_rows = []   # (sta_label, param, best_level, best_mean, worst_mean, corr, trend)

# ── Per-STA-type sheets ─────────────────────────────────────────────────
for sta_type, sta_label in STA_TYPES:
    subset = [r for r in records if r["sta_type"] == sta_type]
    ws = wb.create_sheet(title=sta_label)
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:I1")
    t = ws["A1"]
    t.value = f"{sta_label} — P99 VO delay sensitivity to CWds / QSRC / PSRC  (nSta=30, {DATA_RATE})"
    t.fill, t.font, t.alignment = TITLE_FILL, Font(color="F0F6FC", bold=True, size=12), LEFT
    ws.row_dimensions[1].height = 22

    row = 3
    for param in PARAMS:
        levels = PARAM_LEVELS[param]
        # header
        ws.merge_cells(f"A{row}:I{row}")
        h = ws[f"A{row}"]
        h.value = f"▶ Effect of {param}  (mean P99 µs, averaged over the other 2 params)"
        h.fill, h.font, h.alignment = HEAD_FILL, BLUE, LEFT
        row += 1

        cols = [f"{param}", "mean P99 (all nPedca)", "min P99", "max P99",
                "nPedca=5", "nPedca=15", "nPedca=30", "rows", ""]
        for ci, c in enumerate(cols, 1):
            cell(ws, row, ci, c, BOLD_W, HEAD_FILL, CENTER)
        row += 1

        # compute marginal means
        level_overall = {}
        for lv in levels:
            grp = [r for r in subset if r[param] == lv]
            if not grp:
                continue
            vals = [r[METRIC] for r in grp]
            level_overall[lv] = mean(vals)

        if not level_overall:
            continue
        best_lv = min(level_overall, key=level_overall.get)

        for i, lv in enumerate(levels):
            grp = [r for r in subset if r[param] == lv]
            if not grp:
                continue
            vals = [r[METRIC] for r in grp]
            m = mean(vals)
            per_n = {}
            for n in N_PEDCA_LIST:
                g = [r[METRIC] for r in grp if r["nPedca"] == n]
                per_n[n] = mean(g) if g else None
            bg = BEST_FILL if lv == best_lv else (ZEBRA_A if i % 2 else ZEBRA_B)
            cell(ws, row, 1, lv, BOLD_W, bg, CENTER)
            cell(ws, row, 2, round(m, 1), GREEN if lv == best_lv else WHITE, bg, RIGHT, "#,##0.0")
            cell(ws, row, 3, round(min(vals), 1), WHITE, bg, RIGHT, "#,##0.0")
            cell(ws, row, 4, round(max(vals), 1), WHITE, bg, RIGHT, "#,##0.0")
            cell(ws, row, 5, round(per_n[5], 1) if per_n[5] else None, WHITE, bg, RIGHT, "#,##0.0")
            cell(ws, row, 6, round(per_n[15], 1) if per_n[15] else None, WHITE, bg, RIGHT, "#,##0.0")
            cell(ws, row, 7, round(per_n[30], 1) if per_n[30] else "—", WHITE, bg, RIGHT, "#,##0.0")
            cell(ws, row, 8, len(grp), GREY, bg, CENTER)
            cell(ws, row, 9, "◀ best (min P99)" if lv == best_lv else "", GREEN, bg, LEFT)
            row += 1

        # correlation across all rows
        xs = [r[param] for r in subset]
        ys = [r[METRIC] for r in subset]
        corr = pearson(xs, ys)
        ws.merge_cells(f"A{row}:I{row}")
        c = ws[f"A{row}"]
        rec = ("smaller is better" if level_overall[best_lv] == min(level_overall.values())
               and best_lv == min(level_overall) else "")
        c.value = (f"   → Pearson corr(level, P99) = "
                   f"{corr:+.3f}   |   {trend_text(corr)}   |   "
                   f"recommended {param} = {best_lv} (lowest mean P99 = {level_overall[best_lv]:.0f} µs)")
        c.fill, c.font, c.alignment = ZEBRA_B, GOLD, LEFT
        row += 2

        conclusion_rows.append((sta_label, param, best_lv,
                                level_overall[best_lv],
                                max(level_overall.values()), corr,
                                trend_text(corr)))

    widths = [10, 20, 12, 12, 12, 12, 12, 8, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

# ── Conclusion / cheat-sheet ────────────────────────────────────────────
ws = wb.create_sheet(title="● Conclusion", index=0)
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = "1F6FEB"

ws.merge_cells("A1:G1")
t = ws["A1"]
t.value = "How to tune each parameter to MINIMISE P99 VO delay  (nSta=30, 1Mbps)"
t.fill, t.font, t.alignment = TITLE_FILL, Font(color="F0F6FC", bold=True, size=13), LEFT
ws.row_dimensions[1].height = 24

heads = ["Viewpoint", "Parameter", "Best value (min P99)",
         "Best mean P99 µs", "Worst mean P99 µs", "Pearson corr", "Trend / recommendation"]
for ci, c in enumerate(heads, 1):
    cell(ws, 2, ci, c, BOLD_W, HEAD_FILL, CENTER)

r = 3
for sta_label, param, best_lv, best_m, worst_m, corr, trend in conclusion_rows:
    bg = TYPE_FILL[{"All STAs": "all", "P-EDCA STAs": "pedca", "Legacy STAs": "legacy"}[sta_label]]
    cell(ws, r, 1, sta_label, BOLD_W, bg, LEFT)
    cell(ws, r, 2, param, WHITE, bg, CENTER)
    cell(ws, r, 3, best_lv, GREEN, bg, CENTER)
    cell(ws, r, 4, round(best_m, 1), GREEN, bg, RIGHT, "#,##0.0")
    cell(ws, r, 5, round(worst_m, 1), WHITE, bg, RIGHT, "#,##0.0")
    cell(ws, r, 6, round(corr, 3) if corr is not None else "n/a", WHITE, bg, CENTER, "+0.000;-0.000")
    cell(ws, r, 7, trend, WHITE, bg, LEFT)
    r += 1

# legend
r += 1
ws.merge_cells(f"A{r}:G{r}")
cell(ws, r, 1, "Notes: 'mean P99' is averaged over all other parameters and nPedca. "
               "Pearson corr is between the parameter level and P99 across every combo; "
               "positive = larger param worsens delay (so smaller is better), negative = larger is better. "
               "Legacy STAs only exist at nPedca=5/15 (none at 30).",
     GREY, ZEBRA_B, LEFT)

widths = [16, 12, 20, 18, 18, 14, 42]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A3"

out = BASE / f"param_sensitivity_p99_{DATA_RATE}.xlsx"
wb.save(out)
print(f"✔ wrote {out.name}")

# also print the conclusion to stdout
print("\n=== CONCLUSION (min P99) ===")
for sta_label, param, best_lv, best_m, worst_m, corr, trend in conclusion_rows:
    print(f"{sta_label:13s} {param:5s} best={best_lv}  "
          f"meanP99={best_m:8.0f}us  corr={corr:+.3f}  {trend}")
